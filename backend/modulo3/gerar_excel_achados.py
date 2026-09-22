"""
Excel de saída do Módulo 3 — uma aba "Achados", mesma paleta de cores do
Módulo 2 (modulo2/gerar_excel_final.py), pra manter a identidade visual.
"""
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

COR_HEADER = "4F5C44"       # verde-musgo escuro
COR_TEXTO_HEADER = "FFFFFF"
COR_ALTO = "F5D4CC"         # rosa-chá mais forte — mesma cor de "precisa de atenção" do Módulo 2
COR_MEDIO = "F8EBE8"        # rosa-chá bem claro
COR_BAIXO = "F3E8D0"        # tom quente neutro
COR_OK = "EDF1E7"           # verde-musgo bem claro

_COR_SEVERIDADE = {"Alto": COR_ALTO, "Médio": COR_MEDIO, "Baixo": COR_BAIXO, "OK": COR_OK}
BORDA = Border(*[Side(style="thin", color="E6DFD6")] * 4)

_CABECALHOS = ["Grupo", "Acesso", "Conta", "Terceiro", "Tipo", "Descrição", "Valor (R$)", "Severidade", "Referência"]


def gerar_excel_achados(achados: list, nome_empresa: str = "", periodo: str = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Achados"

    ws["A1"] = f"Análise de razão — {nome_empresa}" + (f" ({periodo})" if periodo else "")
    ws["A1"].font = Font(bold=True, size=12)
    ws.merge_cells("A1:I1")

    linha_header = 3
    for col, texto in enumerate(_CABECALHOS, start=1):
        cell = ws.cell(row=linha_header, column=col, value=texto)
        cell.font = Font(bold=True, color=COR_TEXTO_HEADER)
        cell.fill = PatternFill("solid", fgColor=COR_HEADER)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDA
    ws.row_dimensions[linha_header].height = 26
    ws.freeze_panes = f"A{linha_header + 1}"
    ws.auto_filter.ref = f"A{linha_header}:I{linha_header}"

    for i, a in enumerate(achados, start=linha_header + 1):
        valores = [
            a.grupo.capitalize(), a.acesso, a.nome_conta, a.terceiro or "",
            a.tipo, a.descricao, a.valor, a.severidade, a.referencia or "",
        ]
        cor = _COR_SEVERIDADE.get(a.severidade, "FFFFFF")
        for col, valor in enumerate(valores, start=1):
            cell = ws.cell(row=i, column=col, value=valor)
            cell.fill = PatternFill("solid", fgColor=cor)
            cell.border = BORDA
            cell.alignment = Alignment(vertical="top", wrap_text=(col in (5, 6)))
            if col == 7 and valor is not None:
                cell.number_format = '#,##0.00'

    larguras = [10, 9, 26, 24, 24, 60, 13, 11, 12]
    for col, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(col)].width = largura

    if not achados:
        ws.cell(row=linha_header + 1, column=1, value="Nenhum achado — tudo consistente nas regras aplicadas.")

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()